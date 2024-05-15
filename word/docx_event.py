import openpyxl as excel
import docx
import os

in_file = 'input/name_addr.xlsx'
template_file = 'input/template/event-template.docx'
save_dir = os.path.join(os.path.dirname(__file__), 'output', 'events')

if not os.path.exists(save_dir):
    os.mkdir(save_dir)


def read_book():
    result = []
    sheet = excel.load_workbook(in_file).active

    for row in sheet.iter_rows(min_row=2):
        v = [c.value for c in row]

        if v[0] is None:
            break

        result.append(v)

    return result


for person in read_book():
    (name, zip_num, addr) = person

    card = {
        '[[주소]]': f'(우){zip_num} | {addr}',
        '[[고객명]]': name
    }

    doc = docx.Document(template_file)

    for p in doc.paragraphs:
        for k, v in card.items():
            if p.text.find(k) >= 0:
                p.text = p.text.replace(k, v)
                p.runs[0].font.bold = True

    save_file = os.path.join(save_dir, f'{name} 님.docx')
    print('save: ', save_file)
    doc.save(save_file)
    